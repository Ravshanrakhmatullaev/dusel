#!/usr/bin/env python3
"""
DUSEL GROUP — 10 000 ta QR generator.

Natija (out/ papkasida):
  QR_SVG/00001.svg ... 10000.svg   — vektor, CorelDRAW / bosma uchun (o'lchami config: svg_size_mm)
  QR_PNG/00001.png ... 10000.png   — 1-bit PNG, modul = butun piksel (aniq qirralar)
  bilet_raqamlari.csv / .xlsx      — NUMBER | URL | QR_FILENAME (boshidagi nollar saqlanadi)
  printmerge_6up.csv               — CorelDRAW Print Merge: 1 qator = 1 list (6 ta bilet)
  DUSEL_QR_11222_all.zip               — barcha SVG + PNG + jadvallar

Ishlatish:
  pip install segno pillow openpyxl
  python gen_qr.py [--config config.json] [--out out] [--qr-dir "D:\\DUSEL\\QR_SVG"]
"""
import argparse, csv, json, os, sys, zipfile
from concurrent.futures import ProcessPoolExecutor
import segno
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))


def load_cfg(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def code_str(n, digits):
    return str(n).zfill(digits)


def make_qr(url, q):
    qr = segno.make(url, version=q["version"], error=q["error"].lower(),
                    mode="byte", boost_error=False, micro=False)
    if qr.version != str(q["version"]) and qr.version != q["version"]:
        raise ValueError(f"QR versiya kutilmagan: {qr.version}")
    return qr


def matrix(qr, border):
    return [list(r) for r in qr.matrix_iter(scale=1, border=border)]


def svg_from_matrix(m, size_mm):
    n = len(m)
    d = []
    for y, row in enumerate(m):
        x = 0
        while x < n:
            if row[x]:
                s = x
                while x < n and row[x]:
                    x += 1
                d.append(f"M{s} {y}h{x - s}v1h-{x - s}z")
            else:
                x += 1
    return (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{size_mm}mm" height="{size_mm}mm" '
        f'viewBox="0 0 {n} {n}" shape-rendering="crispEdges">'
        f'<rect width="{n}" height="{n}" fill="#FFFFFF"/>'
        f'<path fill="#000000" d="{"".join(d)}"/></svg>\n'
    )


def png_from_matrix(m, ppm, dpi, path):
    n = len(m)
    img = Image.new("1", (n, n), 1)
    px = img.load()
    for y, row in enumerate(m):
        for x, v in enumerate(row):
            if v:
                px[x, y] = 0
    img = img.resize((n * ppm, n * ppm), Image.NEAREST)
    img.save(path, dpi=(dpi, dpi), optimize=True)


def work(args):
    n, cfg, out = args
    q = cfg["qr"]
    code = code_str(n, cfg["digits"])
    url = cfg["base_url"] + code
    qr = make_qr(url, q)
    m = matrix(qr, q["quiet_zone_modules"])
    with open(os.path.join(out, "QR_SVG", code + ".svg"), "w", encoding="utf-8") as f:
        f.write(svg_from_matrix(m, q["svg_size_mm"]))
    png_from_matrix(m, q["png_px_per_module"], q["png_dpi"], os.path.join(out, "QR_PNG", code + ".png"))
    return code, url


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", default=os.path.join(HERE, "config.json"))
    ap.add_argument("--out", default="out")
    ap.add_argument("--qr-dir", default="", help="Print Merge CSV uchun QR papkaning to'liq yo'li (ixtiyoriy)")
    ap.add_argument("--no-zip", action="store_true")
    a = ap.parse_args()
    cfg = load_cfg(a.config)
    if not cfg["base_url"].startswith("https://"):
        sys.exit("base_url HTTPS bo'lishi shart")
    os.makedirs(os.path.join(a.out, "QR_SVG"), exist_ok=True)
    os.makedirs(os.path.join(a.out, "QR_PNG"), exist_ok=True)

    nums = list(range(cfg["first"], cfg["last"] + 1))
    with ProcessPoolExecutor() as ex:
        rows = list(ex.map(work, [(n, cfg, a.out) for n in nums], chunksize=200))

    codes = [r[0] for r in rows]
    assert len(codes) == len(set(codes)) == len(nums), "takroriy yoki yetishmayotgan raqam!"

    # --- CSV (ASCII/UTF-8, BOMsiz — CorelDRAW Print Merge uchun) ---
    with open(os.path.join(a.out, "bilet_raqamlari.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["NUMBER", "URL", "QR_FILENAME"])
        for code, url in rows:
            w.writerow([code, url, code + ".svg"])

    # --- 6-up Print Merge CSV: 1 qator = 1 list ---
    per = cfg["imposition"]["cols"] * cfg["imposition"]["rows"]
    qdir = a.qr_dir.rstrip("\\/")
    sep = "\\" if ("\\" in qdir or (len(qdir) > 1 and qdir[1] == ":")) else "/"
    with open(os.path.join(a.out, "printmerge_6up.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        hdr = ["SHEET"] + [f"N{i}" for i in range(1, per + 1)] + [f"URL{i}" for i in range(1, per + 1)] \
              + [f"QR{i}" for i in range(1, per + 1)]
        w.writerow(hdr)
        for s in range(0, len(rows), per):
            chunk = rows[s:s + per]
            ns = [c for c, _ in chunk] + [""] * (per - len(chunk))
            us = [u for _, u in chunk] + [""] * (per - len(chunk))
            qs = [((qdir + sep) if qdir else "") + c + ".svg" if c else "" for c in ns]
            w.writerow([s // per + 1] + ns + us + qs)

    # --- XLSX: raqamlar matn formatida ---
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Biletlar"
    ws.append(["NUMBER", "URL", "QR_FILENAME"])
    for code, url in rows:
        ws.append([code, url, code + ".svg"])
    for c in ws["A"]:
        c.number_format = "@"
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="E00010")
    for c in ws[1]:
        c.font, c.fill, c.alignment = hdr_font, hdr_fill, Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 58
    ws.column_dimensions["C"].width = 16
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{len(rows) + 1}"

    ws2 = wb.create_sheet("6-up listlar")
    ws2.append(["SHEET"] + [f"N{i}" for i in range(1, per + 1)] + ["DIAPAZON"])
    for s in range(0, len(rows), per):
        chunk = [c for c, _ in rows[s:s + per]]
        ws2.append([s // per + 1] + chunk + [""] * (per - len(chunk)) + [f"{chunk[0]}–{chunk[-1]}"])
    for col in ws2.iter_cols(min_col=2, max_col=per + 1):
        for c in col:
            c.number_format = "@"
    for c in ws2[1]:
        c.font, c.fill, c.alignment = hdr_font, hdr_fill, Alignment(horizontal="center")
    ws2.freeze_panes = "B2"
    ws2.column_dimensions["H"].width = 16
    wb.save(os.path.join(a.out, "bilet_raqamlari.xlsx"))

    if not a.no_zip:
        zp = os.path.join(a.out, "DUSEL_QR_11222_all.zip")
        with zipfile.ZipFile(zp, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
            for sub in ("QR_SVG", "QR_PNG"):
                for code in codes:
                    ext = sub[-3:].lower()
                    z.write(os.path.join(a.out, sub, f"{code}.{ext}"), f"{sub}/{code}.{ext}")
            for fn in ("bilet_raqamlari.csv", "bilet_raqamlari.xlsx", "printmerge_6up.csv"):
                z.write(os.path.join(a.out, fn), fn)
    print(f"OK: {len(rows)} ta QR, {codes[0]}–{codes[-1]}")


if __name__ == "__main__":
    main()
