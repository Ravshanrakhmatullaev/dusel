#!/usr/bin/env python3
"""места.xlsx o'quvchi: har bir sektor bloki katak joylashuvi bo'yicha (birlashtirilgan kataklar hisobga olinadi)."""
import re
from openpyxl import load_workbook

CYR2LAT = str.maketrans({"А": "A", "В": "B", "С": "C", "Е": "E", "Н": "H", "К": "K", "М": "M", "О": "O",
                         "Р": "P", "Т": "T", "Х": "X"})
ROW_SHEETS = ("блитчеры", "2-этаж", "4-этаж")
VIP_SHEET = "балкон 3-vip"


def _merged_map(ws):
    m = {}
    for rng in ws.merged_cells.ranges:
        v = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                m[(r, c)] = (v, str(rng))
    return m


def _total(v):
    if v is None:
        return None
    s = str(v)
    if "НЕТ" in s.upper():
        return 0
    d = re.sub(r"[^\d]", "", s)
    return int(d) if d else None


def parse(path):
    """-> dict(rows=[...], sectors=[...], floor_stated={sheet: int}, vip_seats=int, sheets=[...])"""
    wb = load_workbook(path)
    rows, sectors, floor_stated = [], [], {}
    vip = 0
    for ws in wb.worksheets:
        sheet = ws.title
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip().startswith("Всего"):
                    floor_stated[sheet] = int(re.sub(r"[^\d]", "", c.value.split("этаже")[1]))
        if sheet == VIP_SHEET:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and re.match(r"^(Ложа|VIP|Президент)", c.value.strip()):
                        v = ws.cell(c.row + 2, c.column).value
                        vip += int(v or 0)
            continue
        mm = _merged_map(ws)
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not (isinstance(v, str) and re.match(r"^(Блитчер|Сектор)\s", v.strip())):
                    continue
                title = v.strip()
                m = re.search(r"([A-Za-zА-Яа-я])\s*(\d{3})", title)
                code = m.group(1) + m.group(2)
                col = c.column
                r = c.row + 1
                while r <= ws.max_row and ws.cell(r, col).value != "Ряд":
                    r += 1
                r += 1
                computed, stated = 0, None
                while r <= ws.max_row:
                    rv = ws.cell(r, col).value
                    if isinstance(rv, str) and rv.strip().startswith("Итого"):
                        stated = _total(rv)
                        break
                    if rv is None:
                        break
                    cnt_raw = ws.cell(r, col + 1).value
                    note_cell = mm.get((r, col + 2))
                    note = note_cell[0] if note_cell else ws.cell(r, col + 2).value
                    note_range = note_cell[1] if note_cell else None
                    special = None
                    if isinstance(cnt_raw, (int, float)):
                        n = int(cnt_raw)
                    elif isinstance(cnt_raw, str):
                        special = cnt_raw.strip()
                        d = re.findall(r"\d+", special)
                        n = int(d[-1]) if d else 0
                    else:
                        n = 0
                    computed += n
                    rows.append(dict(sheet=sheet, sector_original=code, sector=code.translate(CYR2LAT),
                                     sector_num=int(m.group(2)), row=int(rv), count=n, count_raw=cnt_raw,
                                     special=special, note=(note or "").strip(), note_range=note_range,
                                     cell=f"{sheet}!{ws.cell(r, col).coordinate}:{ws.cell(r, col + 2).coordinate}"))
                    r += 1
                sectors.append(dict(sheet=sheet, sector_original=code, sector=code.translate(CYR2LAT),
                                    title=title, stated=stated, computed=computed, header=c.coordinate))
    return dict(rows=rows, sectors=sectors, floor_stated=floor_stated, vip_seats=vip,
                sheets=[w.title for w in wb.worksheets])
